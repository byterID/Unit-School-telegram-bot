"""
Импорт меню столовой из Excel.

    python -m data.import_menu --template меню.xlsx   # шаблон, уже заполненный
    python -m data.import_menu меню.xlsx --dry-run    # проверка
    python -m data.import_menu меню.xlsx              # загрузка

Лист: «День недели | Завтрак | Обед» (можно добавить «Полдник»).
Блюда в ячейке — каждое с новой строки (Alt+Enter) или через «;».
Если «День недели» пуст, строка относится к предыдущему дню.
Меню дней, которые есть в файле, заменяется целиком; остальные дни не трогаются.
"""
from __future__ import annotations

import argparse
import asyncio
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

import config
from database import init_db

EMPTY = {"", "nan", "none", "-", "—", "–"}
SHORT_DAYS = {"пн": 1, "вт": 2, "ср": 3, "чт": 4, "пт": 5, "сб": 6, "вс": 7}
# «🥣 Завтрак» -> «завтрак»: колонки ищем по названиям из config.MEAL_TYPES
MEAL_BY_TITLE = {title.split()[-1].lower(): key for key, title in config.MEAL_TYPES.items()}
BULLET_RE = re.compile(r"^(?:[•·*\-–—]+|\d{1,2}[.)])\s*")

SAMPLE_MENU: dict[int, dict[str, list[str]]] = {
    1: {
        "breakfast": ["Каша овсяная молочная с маслом", "Бутерброд с сыром", "Чай с сахаром"],
        "lunch": ["Салат из свежих огурцов и помидоров", "Суп куриный с вермишелью",
                  "Котлета куриная с картофельным пюре", "Компот из сухофруктов", "Хлеб"],
    },
    2: {
        "breakfast": ["Омлет натуральный", "Хлеб с маслом", "Какао с молоком"],
        "lunch": ["Винегрет", "Борщ со сметаной", "Плов с курицей", "Сок", "Хлеб"],
    },
    3: {
        "breakfast": ["Сырники со сметаной", "Чай с лимоном", "Яблоко"],
        "lunch": ["Салат из капусты с морковью", "Суп гороховый",
                  "Тефтели в соусе с гречкой", "Кисель", "Хлеб"],
    },
    4: {
        "breakfast": ["Каша рисовая молочная", "Бутерброд с маслом и сыром", "Какао"],
        "lunch": ["Салат из свёклы", "Щи из свежей капусты",
                  "Рыба запечённая с рисом", "Компот", "Хлеб"],
    },
    5: {
        "breakfast": ["Запеканка творожная со сгущёнкой", "Чай", "Банан"],
        "lunch": ["Салат «Витаминный»", "Суп-лапша домашняя",
                  "Гуляш с макаронами", "Морс", "Хлеб"],
    },
}


def norm(value) -> str:
    return " ".join(str(value if value is not None else "").split())


def parse_weekday(text: str) -> int | None:
    t = text.lower().rstrip(".")
    for num, name in config.WEEKDAYS.items():
        if t.startswith(name.lower()):
            return num
    return SHORT_DAYS.get(t[:2])


def split_dishes(value) -> list[str]:
    if value is None:
        return []
    result = []
    for part in re.split(r"[\n;]+", str(value).replace("_x000D_", "\n")):
        part = BULLET_RE.sub("", norm(part))
        if part and part.lower() not in EMPTY:
            result.append(part[:200])
    return result


def parse_menu(path: Path) -> tuple[dict[int, dict[str, list[str]]], list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Меню"] if "Меню" in wb.sheetnames else wb.active
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    header_idx = day_col = None
    meal_cols: dict[int, str] = {}
    for h, row in enumerate(rows[:20]):
        cells = [norm(c).lower() for c in row]
        days = [i for i, c in enumerate(cells) if c.startswith("день")]
        meals: dict[int, str] = {}
        for i, c in enumerate(cells):
            for title, key in MEAL_BY_TITLE.items():
                if title in c and key not in meals.values():
                    meals[i] = key
        if days and meals:
            header_idx, day_col, meal_cols = h, days[0], meals
            break
    if header_idx is None:
        raise ValueError("Не найден заголовок «День недели | Завтрак | Обед»")

    menu: dict[int, dict[str, list[str]]] = {}
    warnings: list[str] = []
    weekday: int | None = None
    for n, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        day_text = norm(row[day_col] if day_col < len(row) else None)
        if day_text and day_text.lower() not in EMPTY:
            weekday = parse_weekday(day_text)
            if weekday is None:
                warnings.append(f"строка {n}: не понял день «{day_text}» — пропускаю")
                continue
        for col, meal in meal_cols.items():
            dishes = split_dishes(row[col] if col < len(row) else None)
            if not dishes:
                continue
            if weekday is None:
                warnings.append(f"строка {n}: блюда без дня недели — пропускаю")
                break
            bucket = menu.setdefault(weekday, {}).setdefault(meal, [])
            bucket.extend(d for d in dishes if d not in bucket)
    return menu, warnings


def make_template(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    if path.exists():
        sys.exit(f"Файл {path} уже существует")
    wb = Workbook()
    ws = wb.active
    ws.title = "Меню"
    ws.append(["День недели", "Завтрак", "Обед"])
    for wd, meals in SAMPLE_MENU.items():
        ws.append([config.WEEKDAYS[wd], "\n".join(meals["breakfast"]), "\n".join(meals["lunch"])])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in zip("ABC", (16, 40, 45)):
        ws.column_dimensions[col].width = width
    wb.save(path)
    print(f"Шаблон создан: {path}")


async def save(menu) -> None:
    db = init_db(config.DB_PATH)
    await db.connect()
    try:
        await db.replace_menu(menu)
    finally:
        await db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Импорт меню столовой из Excel")
    parser.add_argument("file", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--template", action="store_true", help="создать заполненный шаблон")
    args = parser.parse_args()

    if args.template:
        make_template(args.file)
        return
    if not args.file.exists():
        sys.exit(f"Файл не найден: {args.file}")

    menu, warnings = parse_menu(args.file)
    if not menu:
        sys.exit("В файле не найдено ни одного блюда")
    for wd in sorted(menu):
        print(f"\n{config.WEEKDAYS[wd]}")
        for meal, dishes in menu[wd].items():
            print(f"  {config.MEAL_TYPES[meal]}: {', '.join(dishes)}")
    for w in warnings:
        print("  ⚠️", w)

    if args.dry_run:
        print("\n--dry-run: в базу ничего не записано.")
        return
    asyncio.run(save(menu))
    print("\n✅ Меню загружено в", config.DB_PATH)


if __name__ == "__main__":
    main()
