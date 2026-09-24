"""
Импорт расписания из Excel (.xlsx) в базу бота.

Запуск из корня проекта:
    python -m data.import_schedule "расписание.xlsx" --dry-run   # только проверка
    python -m data.import_schedule "расписание.xlsx"             # загрузка в БД

Формат листа: «День недели | Урок | Время | 1 класс | 2 класс | ...»
Поддерживаются оба варианта:
    • предмет и ФИО в одной ячейке через перенос строки (Alt+Enter);
    • предмет в одной строке, ФИО — в строке ниже.

Расписание заменяется целиком. Классы и преподаватели сохраняют свои id,
поэтому привязки пользователей и коды приглашения не слетают.
"""
from __future__ import annotations

import argparse
import asyncio
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

import config
from data.import_teachers import compatible
from database import init_db

# Имена, записанные в таблице нестандартно: «как в таблице» -> «полное ФИО».
# Короткие варианты («Мазурова», «Филатов») и переставленные имя/фамилия
# («Эльмира Навасардян») распознаются автоматически.
TEACHER_ALIASES: dict[str, str] = {
    # "Гоар": "Симонянц Гоар Эдуардовна",
}

EMPTY = {"", "nan", "none", "-", "—", "–"}
TIME_RE = re.compile(r"(\d{1,2})[:.](\d{2})\s*[–—-]\s*(\d{1,2})[:.](\d{2})")
NUM_RE = re.compile(r"\d+")
WEEKDAY_BY_NAME = {name.lower(): num for num, name in config.WEEKDAYS.items()}


def clean(value) -> str:
    """Нормализует ячейку: убирает лишние пробелы и «пустые» значения."""
    if value is None:
        return ""
    text = re.sub(r"\s+", " ", str(value)).strip()
    return "" if text.lower() in EMPTY else text


def lines(value) -> list[str]:
    """Ячейка -> непустые строки (Excel иногда хранит перенос как _x000D_)."""
    if value is None:
        return []
    text = str(value).replace("_x000D_", "\n")
    return [c for c in (clean(x) for x in text.splitlines()) if c]


def _cell(row, idx):
    return row[idx] if row is not None and idx is not None and idx < len(row) else None


@dataclass
class Lesson:
    group: str
    weekday: int
    number: int
    start: str
    end: str
    subject: str
    teacher: str | None


@dataclass
class ParseResult:
    groups: list[str] = field(default_factory=list)
    lessons: list[Lesson] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def teachers(self) -> dict[str, str]:
        """ФИО -> основные предметы (для карточки в контактах)."""
        subjects: dict[str, Counter] = defaultdict(Counter)
        for lesson in self.lessons:
            if lesson.teacher:
                subjects[lesson.teacher][lesson.subject] += 1
        return {
            name: ", ".join(s for s, _ in counter.most_common(3))
            for name, counter in sorted(subjects.items())
        }


def _find_header(rows: list[tuple]) -> int:
    for idx, row in enumerate(rows):
        cells = [clean(c).lower() for c in row]
        if any(c.startswith("день") for c in cells) and any("урок" in c for c in cells):
            return idx
    raise ValueError("Не найдена строка заголовка «День недели | Урок | Время | ...»")


def parse_schedule(path: str | Path) -> ParseResult:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Расписание"] if "Расписание" in wb.sheetnames else wb.active
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    result = ParseResult()
    h = _find_header(rows)
    header = [clean(c) for c in rows[h]]
    lower = [c.lower() for c in header]
    day_col = next(i for i, c in enumerate(lower) if c.startswith("день"))
    num_col = next(i for i, c in enumerate(lower) if "урок" in c)
    time_col = next((i for i, c in enumerate(lower) if "время" in c), None)
    if time_col is None:
        raise ValueError("Нет столбца «Время»")
    group_cols = {
        i: name for i, name in enumerate(header)
        if name and i not in (day_col, num_col, time_col)
    }
    result.groups = list(group_cols.values())

    # --- Проход 1: сырые пары (предмет, преподаватель) -----------------------
    raw: list[tuple] = []
    weekday: int | None = None
    i = h + 1
    while i < len(rows):
        row = rows[i]
        num_cell = clean(_cell(row, num_col))
        m = NUM_RE.search(num_cell)
        if not m:
            i += 1
            continue
        number = int(m.group())

        day_cell = clean(_cell(row, day_col)).lower()
        if day_cell:
            if day_cell not in WEEKDAY_BY_NAME:
                raise ValueError(f"Строка {i + 1}: неизвестный день «{day_cell}»")
            weekday = WEEKDAY_BY_NAME[day_cell]
        if weekday is None:
            raise ValueError(f"Строка {i + 1}: не указан день недели")

        t = TIME_RE.search(clean(_cell(row, time_col)))
        if not t:
            result.warnings.append(f"Строка {i + 1}: не распознано время — урок пропущен")
            i += 1
            continue
        start = f"{int(t[1]):02d}:{t[2]}"
        end = f"{int(t[3]):02d}:{t[4]}"

        # Следующая строка без номера урока — это строка с преподавателями
        nxt = rows[i + 1] if i + 1 < len(rows) else None
        if nxt is not None and not clean(_cell(nxt, num_col)):
            teacher_row, step = nxt, 2
        else:
            teacher_row, step = None, 1

        day_name = config.WEEKDAYS[weekday]
        for col, group in group_cols.items():
            parts = lines(_cell(row, col))
            teacher = " ".join(lines(_cell(teacher_row, col)))
            if not parts:
                if teacher:
                    result.warnings.append(
                        f"{group}, {day_name}, {number} урок: преподаватель без предмета"
                    )
                continue
            subject = parts[0]
            if not teacher and len(parts) > 1:
                teacher = " ".join(parts[1:])   # «Предмет⏎ФИО» в одной ячейке
            elif len(parts) > 1:
                subject = " ".join(parts)
            raw.append((group, weekday, number, start, end, subject, teacher or None))
        i += step

    # --- Проход 2: уроки по подгруппам («Информатика Мазурова / Шахматы Пожарский»)
    known_subjects = {r[5] for r in raw}

    def split_prefix(text: str) -> tuple[str, str] | None:
        best = max(
            (k for k in known_subjects if text.startswith(k + " ")), key=len, default=None
        )
        return (best, text[len(best):].strip()) if best else None

    pairs: list[tuple] = []
    for group, wd, num, start, end, subject, teacher in raw:
        s_split = split_prefix(subject)
        t_split = split_prefix(teacher) if teacher else None
        if s_split and t_split:
            pairs.append((group, wd, num, start, end, s_split[0], s_split[1]))
            pairs.append((group, wd, num, start, end, t_split[0], t_split[1]))
        else:
            pairs.append((group, wd, num, start, end, subject, teacher))

    # --- Проход 3: нормализация ФИО -----------------------------------------
    names = {TEACHER_ALIASES.get(p[6], p[6]) for p in pairs if p[6]}
    full_names = {n for n in names if len(n.split()) >= 2}
    resolved: dict[str, str] = {}
    for name in names:
        if name in full_names:
            resolved[name] = name
            continue
        candidates = [f for f in full_names if name in f.split()]
        if len(candidates) == 1:
            resolved[name] = candidates[0]
        else:
            resolved[name] = name
            hint = "несколько совпадений" if candidates else "допишите ФИО в TEACHER_ALIASES"
            result.warnings.append(f"Преподаватель «{name}» записан не полностью ({hint})")

    for group, wd, num, start, end, subject, teacher in pairs:
        if teacher:
            teacher = resolved[TEACHER_ALIASES.get(teacher, teacher)]
        result.lessons.append(Lesson(group, wd, num, start, end, subject, teacher))

    # --- Проход 4: накладки у преподавателей ---------------------------------
    busy: dict[tuple, list[Lesson]] = defaultdict(list)
    for lesson in result.lessons:
        if lesson.teacher:
            busy[(lesson.weekday, lesson.number, lesson.teacher)].append(lesson)
    for (wd, num, teacher), items in sorted(busy.items()):
        if len({x.group for x in items}) < 2:
            continue
        where = ", ".join(f"{x.group} ({x.subject})" for x in items)
        kind = "совместный урок?" if len({x.subject for x in items}) == 1 else "НАКЛАДКА"
        result.warnings.append(
            f"{kind}: {teacher} — {config.WEEKDAYS[wd]}, {num} урок: {where}"
        )
    return result


# --- Сопоставление с педагогами в базе ---------------------------------------
def _words(name: str) -> list[str]:
    return name.lower().replace("ё", "е").replace(".", " ").split()


def same_person(a: str, b: str) -> bool:
    """«Навасардян Эльмира» = «Эльмира Навасардян»; «Филатов» = «Филатов Артем»."""
    wa, wb = _words(a), _words(b)
    if not wa or not wb:
        return False
    return sorted(wa) == sorted(wb) or compatible(wa, wb)


async def save(result: ParseResult) -> list[str]:
    notes: list[str] = []
    db = init_db(config.DB_PATH)
    await db.connect()
    try:
        group_ids: dict[str, int] = {}
        for group in result.groups:
            m = NUM_RE.search(group)
            group_ids[group] = await db.upsert_group(group, int(m.group()) if m else 999)

        existing = [dict(t) for t in await db.get_teachers()]
        teacher_ids: dict[str, int] = {}
        for name, subjects in result.teachers.items():
            matches = [t for t in existing if same_person(t["full_name"], name)]

            # «Абрамян» подходит и к Жанне, и к Элине — это разные люди, не сливаем
            if len(matches) > 1 and not all(
                same_person(matches[0]["full_name"], t["full_name"]) for t in matches[1:]
            ):
                exact = [t for t in matches if t["full_name"] == name]
                if not exact:
                    notes.append(
                        f"⚠️ «{name}» подходит нескольким: "
                        + ", ".join(t["full_name"] for t in matches)
                        + " — уроки останутся без преподавателя"
                    )
                    continue
                matches = exact

            if not matches:
                tid = await db.upsert_teacher(name, subjects)
                existing.append({"id": tid, "full_name": name})
                notes.append(f"+ новый педагог: {name}")
                teacher_ids[name] = tid
                continue

            # Оставляем запись с точным совпадением, иначе — с самым длинным ФИО
            matches.sort(key=lambda t: (t["full_name"] != name, -len(t["full_name"])))
            keep = matches[0]
            for dup in matches[1:]:
                await db.merge_teachers(keep["id"], dup["id"])
                existing = [t for t in existing if t["id"] != dup["id"]]
                notes.append(f"~ дубль «{dup['full_name']}» объединён с «{keep['full_name']}»")
            new_name = name if len(name) > len(keep["full_name"]) else keep["full_name"]
            await db.update_teacher(keep["id"], new_name, subjects)
            keep["full_name"] = new_name
            teacher_ids[name] = keep["id"]

        await db.replace_schedule(
            (
                group_ids[l.group],
                teacher_ids.get(l.teacher) if l.teacher else None,
                l.weekday, l.number, l.start, l.end, l.subject, None,
            )
            for l in result.lessons
        )
    finally:
        await db.close()
    return notes


def main() -> None:
    parser = argparse.ArgumentParser(description="Импорт расписания из Excel")
    parser.add_argument("file", type=Path, help="Путь к .xlsx")
    parser.add_argument("--dry-run", action="store_true", help="Только проверить, не записывать")
    args = parser.parse_args()
    if not args.file.exists():
        sys.exit(f"Файл не найден: {args.file}")

    result = parse_schedule(args.file)
    per_teacher = Counter(l.teacher for l in result.lessons if l.teacher)
    no_teacher = [l for l in result.lessons if not l.teacher]

    print(f"Классов: {len(result.groups)}")
    print(f"Уроков: {len(result.lessons)}  (без преподавателя: {len(no_teacher)})")
    print(f"Преподавателей: {len(per_teacher)}")
    for name, count in sorted(per_teacher.items()):
        print(f"  {name}: {count} ур.")
    for l in no_teacher[:10]:
        print(
            f"  ⚠️ без преподавателя: {l.group}, {config.WEEKDAYS[l.weekday]}, "
            f"{l.number} урок — «{l.subject}»"
        )
    if result.warnings:
        print(f"\nПредупреждения ({len(result.warnings)}):")
        for w in result.warnings:
            print("  •", w)

    if args.dry_run:
        print("\n--dry-run: в базу ничего не записано.")
        return
    for note in asyncio.run(save(result)):
        print(note)
    print("\n✅ Расписание загружено в", config.DB_PATH)


if __name__ == "__main__":
    main()
