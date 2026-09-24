"""
FAQ: список по умолчанию и импорт из Excel.

    python -m data.import_faq                       # загрузить список ниже (заменит текущий)
    python -m data.import_faq --template faq.xlsx   # выгрузить его в Excel для правки
    python -m data.import_faq faq.xlsx              # загрузить из Excel («Вопрос | Ответ»)
"""
from __future__ import annotations

import argparse
import asyncio
import sys
from pathlib import Path

import config
from database import init_db

DEFAULT_FAQ: list[tuple[str, str]] = [
    ("Во сколько начинаются уроки?",
     "Первый урок — в 08:45. У старших классов занятия обычно начинаются со 2–4 урока. "
     "Точное время на нужный день — в «📅 Расписание» → «⭐️ Мой класс»."),
    ("Какое расписание звонков?",
     "1 урок — 08:45–09:25\n2 урок — 09:40–10:20\n3 урок — 10:30–11:10\n"
     "4 урок — 11:20–12:00\n5 урок — 12:15–12:55\n6 урок — 13:00–13:40\n"
     "7 урок — 13:45–14:25\n8 урок — 14:40–15:20"),
    ("До скольки идут занятия?",
     "1–4 классы обычно заканчивают в 12:00 (по вторникам у 2–4 классов — в 12:55), "
     "5 класс — в 12:55–13:40, 6–10 классы — в 13:40–14:25, у 11 класса занятия до 15:20. "
     "Точно на нужный день — в «📅 Расписание»."),
    ("Как посмотреть расписание своего класса?",
     "«📅 Расписание» → «⭐️ Мой класс», затем «Сегодня», «Завтра», день недели или вся неделя. "
     "Расписание другого класса — через «👥 Выбрать класс»."),
    ("Как узнать о заменах и отменах уроков?",
     "Изменения сразу видны в расписании. Когда администрация вносит замену, бот присылает "
     "уведомление классу — если объявления включены в «⚙️ Настройки»."),
    ("Как посмотреть расписание учителя?",
     "«📅 Расписание» → «👨‍🏫 По преподавателю». Или откройте карточку учителя "
     "в «📞 Контакты» и нажмите «📅 Расписание»."),
    ("Где посмотреть меню столовой?",
     "Кнопка «🍽 Столовая»: завтрак и обед на сегодня, завтра, любой день или всю неделю."),
    ("Как связаться с учителем?",
     "Раздел «📞 Контакты»: выберите преподавателя — там предмет и контакты, которые указала школа."),
    ("Где взять код для входа в бот?",
     "Код класса или ссылку-приглашение выдаёт классный руководитель или администрация школы. "
     "Код можно просто отправить боту сообщением."),
    ("Код не подходит — что делать?",
     "Проверьте, что код набран без ошибок (регистр неважен). Если код перевыпускали, старый "
     "перестаёт работать — попросите новый у классного руководителя. После 5 неверных попыток "
     "ввод блокируется на 10 минут."),
    ("У меня дети в разных классах",
     "В «⚙️ Настройки» → «🎒 Сменить класс» введите код другого класса — он станет «моим». "
     "Расписание любого класса можно открыть через «📅 Расписание» → «👥 Выбрать класс»."),
    ("Как отключить уведомления?",
     "«⚙️ Настройки» → «🔕 Отключить объявления». Включить обратно можно там же."),
    ("Почему мне не приходят объявления?",
     "Проверьте, что объявления включены в «⚙️ Настройки». Если вы останавливали бота, "
     "напишите ему «привет» — доставка возобновится."),
    ("Какие есть дополнительные занятия?",
     "По расписанию проходят: шахматы (1–8 классы), подготовка к ОГЭ (9 класс) и к ЕГЭ "
     "(11 класс), курс психологии «Понять себя, понять других» (10–11 классы). "
     "Дни и время — в «📅 Расписание»."),
    ("Как сообщить, что ребёнок не придёт?",
     "Предупредите классного руководителя заранее — лучше до начала первого урока (08:45)."),
    ("Бот показывает что-то не то",
     f"Напишите администратору {config.SUPPORT_CONTACT} — опишите, что и в каком разделе не так."),
]


def read_xlsx(path: Path) -> list[tuple[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows = [[("" if v is None else str(v)).strip() for v in r]
            for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    for h, row in enumerate(rows[:10]):
        low = [c.lower() for c in row]
        q = next((i for i, c in enumerate(low) if c.startswith("вопрос")), None)
        a = next((i for i, c in enumerate(low) if c.startswith("ответ")), None)
        if q is not None and a is not None:
            break
    else:
        sys.exit("Не найден заголовок «Вопрос | Ответ»")
    items = []
    for row in rows[h + 1:]:
        question = row[q] if q < len(row) else ""
        answer = row[a] if a < len(row) else ""
        if question and answer:
            items.append((question[:200], answer[:3500]))
    return items


def make_template(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    if path.exists():
        sys.exit(f"Файл {path} уже существует")
    wb = Workbook()
    ws = wb.active
    ws.title = "FAQ"
    ws.append(["Вопрос", "Ответ"])
    for item in DEFAULT_FAQ:
        ws.append(list(item))
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 90
    wb.save(path)
    print(f"Шаблон создан: {path}")


async def save(items) -> None:
    db = init_db(config.DB_PATH)
    await db.connect()
    try:
        await db.replace_faq(items)
    finally:
        await db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Загрузка FAQ")
    parser.add_argument("file", type=Path, nargs="?")
    parser.add_argument("--template", action="store_true")
    args = parser.parse_args()

    if args.template:
        if not args.file:
            sys.exit("Укажите путь: --template faq.xlsx")
        make_template(args.file)
        return
    items = read_xlsx(args.file) if args.file else DEFAULT_FAQ
    if not items:
        sys.exit("Нет ни одного вопроса")
    asyncio.run(save(items))
    print(f"✅ Загружено вопросов: {len(items)}")


if __name__ == "__main__":
    main()
