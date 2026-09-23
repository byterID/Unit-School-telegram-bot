"""
Импорт списка педагогов из Excel (.xlsx) или CSV.

Первая строка — заголовки, порядок колонок любой.
Обязательная колонка: ФИО. Необязательные: Предмет, Кабинет, Телефон, Email.

Педагоги, уже созданные импортом расписания («Стяжкина», «Стяжкина И.В.»),
находятся по фамилии и инициалам и дополняются, а не дублируются.

Запуск:
    python -m data.import_teachers --template педагоги.xlsx   # создать шаблон
    python -m data.import_teachers педагоги.xlsx --dry-run    # проверка без записи
    python -m data.import_teachers педагоги.xlsx              # импорт
"""
from __future__ import annotations

import argparse
import csv
import os
import re
import sqlite3
import sys
from pathlib import Path

from dotenv import load_dotenv

# Поле в БД -> с чего может начинаться заголовок колонки
HEADERS: dict[str, tuple[str, ...]] = {
    "full_name": ("фио", "педагог", "преподаватель", "учитель", "сотрудник"),
    "subject": ("предмет", "дисциплин"),
    "room": ("кабинет", "каб"),
    "phone": ("телефон", "тел", "моб"),
    "email": ("email", "e-mail", "почта", "эл"),
}
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


# --- Чтение файла ------------------------------------------------------------
def cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # телефон, сохранённый в Excel числом
    return " ".join(str(value).split())


def read_rows(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        rows = [[cell(v) for v in row] for row in wb.active.iter_rows(values_only=True)]
        wb.close()
        return rows
    if suffix == ".csv":
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            sys.exit("Не удалось определить кодировку CSV (нужна UTF-8 или Windows-1251)")
        delimiter = ";" if text.count(";") >= text.count(",") else ","
        return [[cell(v) for v in row] for row in csv.reader(text.splitlines(), delimiter=delimiter)]
    sys.exit("Поддерживаются только .xlsx и .csv")


def map_header(row: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, title in enumerate(row):
        t = title.lower().replace("ё", "е").replace(".", "").strip()
        for field, prefixes in HEADERS.items():
            if field not in mapping and t.startswith(prefixes):
                mapping[field] = idx
                break
    return mapping


# --- Нормализация ------------------------------------------------------------
def name_key(name: str) -> list[str]:
    """'Стяжкина И.В.' -> ['стяжкина', 'и', 'в']"""
    return name.lower().replace("ё", "е").replace(".", " ").split()


def compatible(a: list[str], b: list[str]) -> bool:
    """Совпадает ли фамилия, а имя/отчество — полностью или по инициалам."""
    if not a or not b or a[0] != b[0]:
        return False
    if len(a) > len(b):
        a, b = b, a
    for x, y in zip(a[1:], b[1:]):
        if not (x == y or (len(x) == 1 and y.startswith(x)) or (len(y) == 1 and x.startswith(y))):
            return False
    return True


def norm_phone(raw: str) -> str | None:
    digits = re.sub(r"\D", "", raw)
    if len(digits) == 11 and digits[0] in "78":
        return "+7" + digits[1:]
    if len(digits) == 10 and digits[0] == "9":
        return "+7" + digits
    return None


def parse(rows: list[list[str]], warnings: list[str]) -> list[dict]:
    for header_idx, row in enumerate(rows[:10]):
        mapping = map_header(row)
        if "full_name" in mapping:
            break
    else:
        sys.exit("Не нашёл строку заголовков с колонкой «ФИО» в первых 10 строках")

    items: list[dict] = []
    seen: dict[str, int] = {}
    for n, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):

        def get(field: str) -> str:
            i = mapping.get(field)
            return row[i] if i is not None and i < len(row) else ""

        name = get("full_name")
        if not name:
            continue
        key = " ".join(name_key(name))
        if key in seen:
            warnings.append(f"строка {n}: «{name}» уже было в строке {seen[key]}, пропускаю")
            continue
        seen[key] = n

        item: dict = {"_row": n, "full_name": name}
        for field in ("subject", "room"):
            if value := get(field):
                item[field] = value
        if phone := get("phone"):
            normalized = norm_phone(phone)
            if normalized is None:
                warnings.append(f"строка {n}: телефон «{phone}» не распознан, оставляю как есть")
            item["phone"] = normalized or phone
        if email := get("email"):
            if EMAIL_RE.match(email):
                item["email"] = email.lower()
            else:
                warnings.append(f"строка {n}: email «{email}» некорректен, пропускаю его")
        items.append(item)
    return items


# --- Запись в БД -------------------------------------------------------------
def apply(conn: sqlite3.Connection, items: list[dict], dry: bool, warnings: list[str]) -> None:
    cols = {r[1] for r in conn.execute("PRAGMA table_info(teachers)")}
    if "full_name" not in cols:
        sys.exit("В таблице teachers нет колонки full_name — пришлите database.py, адаптирую скрипт")

    # Колонки берутся только из белого списка HEADERS, поэтому f-строки безопасны
    missing = {f for it in items for f in it if not f.startswith("_")} - cols
    for col in sorted(missing):
        print(f"+ в таблицу teachers {'будет добавлена' if dry else 'добавлена'} колонка {col}")
        if not dry:
            conn.execute(f"ALTER TABLE teachers ADD COLUMN {col} TEXT")

    existing = [(r[0], r[1]) for r in conn.execute("SELECT id, full_name FROM teachers")]
    used: dict[int, int] = {}
    added = updated = skipped = 0

    for it in items:
        fields = {k: v for k, v in it.items() if not k.startswith("_")}
        new_key = name_key(it["full_name"])
        cands = [(tid, nm) for tid, nm in existing if compatible(name_key(nm), new_key)]

        if len(cands) > 1:
            names = ", ".join(nm for _, nm in cands)
            warnings.append(f"строка {it['_row']}: «{it['full_name']}» подходит к нескольким: {names} — пропускаю")
            skipped += 1
            continue

        if cands:
            tid, old = cands[0]
            if tid in used:
                warnings.append(
                    f"строка {it['_row']}: «{it['full_name']}» совпала с тем же педагогом, "
                    f"что и строка {used[tid]} — пропускаю"
                )
                skipped += 1
                continue
            used[tid] = it["_row"]
            if len(" ".join(name_key(old))) >= len(" ".join(new_key)):
                fields.pop("full_name")  # не заменяем полное ФИО более коротким
            if fields and not dry:
                sets = ", ".join(f"{k} = ?" for k in fields)
                conn.execute(f"UPDATE teachers SET {sets} WHERE id = ?", (*fields.values(), tid))
            print(f"~ {old}: {', '.join(fields) or 'без изменений'}")
            updated += 1
        else:
            if not dry:
                names = ", ".join(fields)
                marks = ", ".join("?" * len(fields))
                conn.execute(f"INSERT INTO teachers ({names}) VALUES ({marks})", tuple(fields.values()))
            print(f"+ {it['full_name']}")
            added += 1

    print(f"\nДобавлено: {added}, обновлено: {updated}, пропущено: {skipped}")


def make_template(path: Path) -> None:
    from openpyxl import Workbook

    if path.exists():
        sys.exit(f"Файл {path} уже существует")
    wb = Workbook()
    ws = wb.active
    ws.title = "Педагоги"
    ws.append(["ФИО", "Предмет", "Кабинет", "Телефон", "Email"])
    ws.append(["Иванова Анна Петровна", "Физика", "214", "89001234567", "fiz@school.ru"])
    ws.append(["Петров Олег Иванович", "Математика", "301", "", ""])
    for col, width in zip("ABCDE", (36, 20, 10, 16, 26)):
        ws.column_dimensions[col].width = width
    wb.save(path)
    print(f"Шаблон создан: {path}")


def main() -> None:
    load_dotenv()
    parser = argparse.ArgumentParser(description="Импорт педагогов")
    parser.add_argument("file", type=Path)
    parser.add_argument("--dry-run", action="store_true", help="только проверить, ничего не записывать")
    parser.add_argument("--template", action="store_true", help="создать шаблон таблицы по пути file")
    parser.add_argument("--db", default=os.getenv("DB_PATH", "school_bot.db"))
    args = parser.parse_args()

    if args.template:
        make_template(args.file)
        return
    if not args.file.exists():
        sys.exit(f"Файл не найден: {args.file}")
    if not Path(args.db).exists():
        sys.exit(f"База {args.db} не найдена — сначала запустите бота или импорт расписания")

    warnings: list[str] = []
    items = parse(read_rows(args.file), warnings)
    if not items:
        sys.exit("В файле не найдено ни одного педагога")

    conn = sqlite3.connect(args.db, timeout=10)
    try:
        with conn:  # одна транзакция: либо всё, либо ничего
            apply(conn, items, args.dry_run, warnings)
            if args.dry_run:
                conn.rollback()
    finally:
        conn.close()

    if warnings:
        print("\n⚠️ Предупреждения:")
        for w in warnings:
            print("  •", w)
    if args.dry_run:
        print("\nЭто была проверка (--dry-run), база не изменена.")


if __name__ == "__main__":
    main()
